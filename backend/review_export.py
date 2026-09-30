from __future__ import annotations

from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


HEADERS = (
    '序号','单词','音标','中文释义','录入日期','最后复习日期',
    '练习次数','正确次数','正确率','掌握度','下一次复习日期','分类',
)


def _excel_date(value):
    text=str(value or '').strip()
    if not text:
        return None
    try:
        return date.fromisoformat(text[:10])
    except ValueError:
        return None


def _safe_text(value):
    text=str(value or '')
    # Prevent spreadsheet applications from interpreting imported vocabulary as a formula.
    return "'"+text if text.startswith(('=','+','-','@')) else text


def daily_words_workbook(words,day,scope_label):
    """Create the exact ordered word list stored in one review session."""
    if not words:
        raise ValueError('本场复习没有可导出的单词')
    book=Workbook()
    sheet=book.active
    sheet.title='当天单词'
    sheet.sheet_view.showGridLines=False
    book.properties.title=f'{day} 当天单词记录'
    book.properties.subject='WordLearner 一键复习清单'
    book.properties.creator='WordLearner'

    sheet.merge_cells('A1:L1')
    sheet['A1']=f'当天单词记录 {day.replace("-","")}'
    sheet['A1'].font=Font(name='Microsoft YaHei',size=18,bold=True,color='FFFFFF')
    sheet['A1'].fill=PatternFill('solid',fgColor='1F4E78')
    sheet['A1'].alignment=Alignment(horizontal='center',vertical='center')
    sheet.row_dimensions[1].height=32

    metadata=(('A2','C2',f'学习日期：{day}'),('D2','F2',f'复习范围：{scope_label}'),
              ('G2','I2',f'单词数量：{len(words)}'),('J2','L2','数据来源：本场实际复习清单'))
    for start,end,value in metadata:
        sheet.merge_cells(f'{start}:{end}')
        cell=sheet[start]; cell.value=value
        cell.font=Font(name='Microsoft YaHei',size=10,color='44546A')
        cell.fill=PatternFill('solid',fgColor='D9EAF7')
        cell.alignment=Alignment(horizontal='left',vertical='center')
    sheet.row_dimensions[2].height=24

    for column,header in enumerate(HEADERS,1):
        cell=sheet.cell(4,column,header)
        cell.font=Font(name='Microsoft YaHei',size=10,bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='2F75B5')
        cell.alignment=Alignment(horizontal='center',vertical='center')
    sheet.row_dimensions[4].height=25

    thin=Side(style='thin',color='D9E2F3')
    for index,item in enumerate(words,1):
        practice=int(item.get('practice_count') or 0)
        correct=int(item.get('correct_count') or 0)
        accuracy=correct/practice if practice else 0
        values=(
            index,_safe_text(item.get('word')),_safe_text(item.get('phonetic')),
            _safe_text(item.get('translation')),_excel_date(item.get('created_on')),
            _excel_date(item.get('last_practiced_at')),practice,correct,accuracy,
            (int(item.get('mastery') or 0)/100),_excel_date(item.get('due_on')),
            '、'.join(item.get('tag_labels') or []),
        )
        row=index+4
        for column,value in enumerate(values,1):
            cell=sheet.cell(row,column,value)
            cell.font=Font(name='Microsoft YaHei',size=10)
            cell.alignment=Alignment(vertical='center',wrap_text=column in (4,12),
                                     horizontal='center' if column in (1,3,5,6,7,8,9,10,11) else 'left')
            cell.border=Border(bottom=thin)
        for column in (5,6,11):
            sheet.cell(row,column).number_format='yyyy-mm-dd'
        for column in (9,10):
            sheet.cell(row,column).number_format='0%'
        sheet.row_dimensions[row].height=24

    last_row=len(words)+4
    table=Table(displayName='DailyReviewWords',ref=f'A4:L{last_row}')
    table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,
                                       showLastColumn=False,showRowStripes=True,showColumnStripes=False)
    sheet.add_table(table)
    widths=(7,24,18,42,13,15,11,11,10,10,15,16)
    for column,width in enumerate(widths,1):
        sheet.column_dimensions[chr(64+column)].width=width
    sheet.freeze_panes='A5'
    sheet.auto_filter.ref=f'A4:L{last_row}'
    sheet.print_title_rows='1:4'
    sheet.page_setup.orientation='landscape'
    sheet.page_setup.fitToWidth=1
    sheet.sheet_properties.pageSetUpPr.fitToPage=True
    sheet.print_options.horizontalCentered=True

    output=BytesIO()
    book.save(output)
    return output.getvalue()
