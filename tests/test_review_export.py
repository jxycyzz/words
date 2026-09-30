from io import BytesIO

from openpyxl import load_workbook

from backend.review_export import HEADERS, daily_words_workbook


def test_daily_workbook_keeps_exact_order_and_learning_fields():
    words=[]
    for index in range(75):
        words.append({
            'word': '=formula' if index == 0 else f'word-{index:02}',
            'phonetic': f'/w{index}/',
            'translation': f'释义 {index}',
            'created_on': '2026-09-01',
            'last_practiced_at': '2026-09-30T08:30:00+08:00' if index else None,
            'practice_count': index,
            'correct_count': index // 2,
            'mastery': 65,
            'due_on': '2026-10-01',
            'tag_labels': ['初二上'],
        })

    content=daily_words_workbook(words,'2026-09-30','初二上')
    sheet=load_workbook(BytesIO(content),data_only=False).active

    assert sheet.title=='当天单词'
    assert sheet['A1'].value=='当天单词记录 20260930'
    assert sheet['G2'].value=='单词数量：75'
    assert tuple(cell.value for cell in sheet[4])==HEADERS
    assert sheet.max_row==79 and sheet.max_column==12
    assert sheet['A5'].value==1 and sheet['B5'].value=="'=formula"
    assert sheet['B79'].value=='word-74'
    assert sheet['I15'].value==5/10 and sheet['I15'].number_format=='0%'
    assert sheet['J5'].value==.65 and sheet['J5'].number_format=='0%'
    assert sheet.freeze_panes=='A5'
    assert 'DailyReviewWords' in sheet.tables

