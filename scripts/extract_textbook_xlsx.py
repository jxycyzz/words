"""Extract the reviewed textbook summary sheet into a deduplicated import list."""
from __future__ import annotations

import argparse
import json
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


def canonical(value):
    value=unicodedata.normalize('NFKC',str(value or ''))
    value=value.translate(str.maketrans({'’':"'",'‘':"'",'`':"'"}))
    return ' '.join(value.strip().split()).casefold()


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('source')
    parser.add_argument('output')
    parser.add_argument('--sheet',default='单词总表')
    args=parser.parse_args()
    workbook=load_workbook(args.source,read_only=True,data_only=True)
    if args.sheet not in workbook.sheetnames:
        raise SystemExit(f'找不到工作表：{args.sheet}')
    rows=[]
    duplicates=[]
    seen=set()
    for row in workbook[args.sheet].iter_rows(min_row=3,values_only=True):
        if len(row)<6 or not row[2]:
            continue
        key=canonical(row[2])
        if key in seen:
            duplicates.append(str(row[2]).strip())
            continue
        seen.add(key)
        rows.append({'unit':str(row[0] or '').strip(),'word':str(row[2]).strip(),
                     'phonetic':str(row[3] or '').strip(),'translation':str(row[5] or '').strip()})
    output=Path(args.output)
    output.parent.mkdir(parents=True,exist_ok=True)
    output.write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps({'rows':len(rows),'duplicates':duplicates,'output':str(output)},ensure_ascii=False))


if __name__=='__main__':
    main()
